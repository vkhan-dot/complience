import os
import re
import json

# Optional heavy parser
try:
    from docling.document_converter import DocumentConverter
    HAS_DOCLING = True
except ImportError:
    HAS_DOCLING = False

# Optional fast HTML parser backend for BeautifulSoup
try:
    import lxml
    HAS_LXML = True
except ImportError:
    HAS_LXML = False

# Fallbacks
from bs4 import BeautifulSoup
import docx
import pypdf

class ParsingAgent:
    """
    ParsingAgent converts documents (HTML, PDF, DOCX) into clean Markdown.
    It attempts to use Docling first, and falls back to lighter standard python libraries if unavailable.
    """
    
    @classmethod
    def parse_file(cls, file_path: str) -> tuple[str, list[dict]]:
        """
        Parses a local file and returns (markdown_content, section_structure).
        """
        ext = os.path.splitext(file_path)[1].lower()
        
        # If Docling is installed, try using it first
        if HAS_DOCLING:
            try:
                converter = DocumentConverter()
                result = converter.convert(file_path)
                md_content = result.document.export_to_markdown()
                structure = cls._extract_sections(md_content)
                return md_content, structure
            except Exception as e:
                # Fallback to standard libraries on error
                print(f"Docling conversion failed: {e}. Falling back to standard parsers.")
        
        # Standard Fallbacks
        if ext == ".html" or ext == ".htm":
            return cls._parse_html_file(file_path)
        elif ext == ".docx":
            return cls._parse_docx_file(file_path)
        elif ext == ".pdf":
            return cls._parse_pdf_file(file_path)
        elif ext == ".txt" or ext == ".md":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return content, cls._extract_sections(content)
        else:
            # General text read
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                return content, cls._extract_sections(content)
            except Exception as e:
                raise ValueError(f"Unsupported file format: {ext} and error reading: {e}")

    @classmethod
    def parse_content(cls, content_bytes: bytes, file_name: str) -> tuple[str, list[dict]]:
        """
        Parses binary content and returns (markdown_content, section_structure).
        """
        import uuid
        ext = os.path.splitext(file_name)[1].lower()
        temp_path = f"temp_parse_{uuid.uuid4().hex}{ext}"
        try:
            with open(temp_path, "wb") as f:
                f.write(content_bytes)
            return cls.parse_file(temp_path)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    @classmethod
    def _clean_html_soup(cls, soup: BeautifulSoup) -> BeautifulSoup:
        """Removes navigational and clutter elements from HTML."""
        unwanted_tags = [
            "script", "style", "header", "footer", "nav", "noscript", "aside", 
            "form", "iframe", "button", "svg", "canvas"
        ]
        for tag in soup(unwanted_tags):
            tag.decompose()

        # Remove portal-specific ads and widgets
        unwanted_selectors = [
            ".doc-tools", ".doc-header", ".doc-footer", ".banner", ".advertisement",
            ".comment-block", ".social-share", ".doc_print", ".doc_nav", ".doc_lang",
            ".sidebar", ".cookie-banner", ".auth-modal", ".popover", ".tooltip"
        ]
        for sel in unwanted_selectors:
            for el in soup.select(sel):
                el.decompose()

        return soup

    @classmethod
    def _parse_html_file(cls, file_path: str) -> tuple[str, list[dict]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            raw_html = f.read()

        soup = BeautifulSoup(raw_html, "lxml" if HAS_LXML else "html.parser")
        soup = cls._clean_html_soup(soup)

        # Target specific legal document containers if present
        doc_container = (
            soup.find("main", class_="document-page") or
            soup.find("div", class_="document-content") or
            soup.find("div", id="doc-body") or
            soup.find("div", class_="doc-text") or
            soup.find("div", class_="content-wrap") or
            soup.find("div", class_="block_doc") or
            soup.find("div", id="law_text") or
            soup.find("div", class_="document_text") or
            soup.find("div", class_="document") or
            soup.find("div", id="document") or
            soup.find("article") or
            soup.find("main")
        )
        target_root = doc_container if (doc_container and len(doc_container.get_text().strip()) > 200) else soup

        markdown_lines = []
        blocks = target_root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "code", "table", "blockquote", "div"])

        # Track tables to avoid double-processing internal paragraphs/cells
        processed_table_descendants = set()
        for tbl in target_root.find_all("table"):
            for desc in tbl.descendants:
                processed_table_descendants.add(id(desc))

        article_regex = re.compile(r'^(статья\s+\d+[\-\d]*|глава\s+\d+|раздел\s+\d+|параграф\s+\d+)\.?\s*(.*)$', re.IGNORECASE)

        for element in blocks:
            if id(element) in processed_table_descendants and element.name != "table":
                continue

            if element.name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
                level = int(element.name[1])
                text = element.get_text().strip()
                if text:
                    markdown_lines.append(f"\n{'#' * level} {text}\n")
            elif element.name == "p":
                text = element.get_text().strip()
                if text:
                    # Check if paragraph is an article or chapter header
                    art_match = article_regex.match(text)
                    if art_match:
                        markdown_lines.append(f"\n### {text}\n")
                    else:
                        markdown_lines.append(f"\n{text}\n")
            elif element.name == "div" and len(element.find_all(["p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "table"])) == 0:
                text = element.get_text().strip()
                if text and len(text) > 2:
                    art_match = article_regex.match(text)
                    if art_match:
                        markdown_lines.append(f"\n### {text}\n")
                    else:
                        markdown_lines.append(f"\n{text}\n")
            elif element.name == "li":
                text = element.get_text().strip()
                if text:
                    markdown_lines.append(f"- {text}")
            elif element.name == "blockquote":
                text = element.get_text().strip()
                if text:
                    markdown_lines.append(f"\n> {text}\n")
            elif element.name in ["pre", "code"] and element.parent.name != "pre":
                text = element.get_text().strip()
                if text:
                    markdown_lines.append(f"\n```\n{text}\n```\n")
            elif element.name == "table":
                rows = element.find_all("tr")
                if rows:
                    markdown_lines.append("\n")
                    for i, row in enumerate(rows):
                        cells = row.find_all(["th", "td"])
                        row_cells = [c.get_text().strip().replace("\n", " ") for c in cells]
                        if row_cells:
                            markdown_lines.append("| " + " | ".join(row_cells) + " |")
                            if i == 0:
                                markdown_lines.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
                    markdown_lines.append("\n")

        # Fallback to plain text if markdown parser was empty
        if not markdown_lines:
            text = soup.get_text()
            chunks = (phrase.strip() for line in text.splitlines() for phrase in line.split("  "))
            content = '\n'.join(chunk for chunk in chunks if chunk)
        else:
            content = "\n".join(markdown_lines)

        content = re.sub(r'\n\s*\n', '\n\n', content)
        return content, cls._extract_sections(content)


    @classmethod
    def _parse_docx_file(cls, file_path: str) -> tuple[str, list[dict]]:
        doc = docx.Document(file_path)
        markdown_lines = []
        article_regex = re.compile(r'^(статья\s+\d+[\-\d]*|глава\s+\d+|раздел\s+\d+)\.?\s*(.*)$', re.IGNORECASE)

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name.lower()
            if "heading 1" in style_name:
                markdown_lines.append(f"\n# {text}\n")
            elif "heading 2" in style_name:
                markdown_lines.append(f"\n## {text}\n")
            elif "heading 3" in style_name:
                markdown_lines.append(f"\n### {text}\n")
            elif "list bullet" in style_name:
                markdown_lines.append(f"- {text}")
            else:
                art_match = article_regex.match(text)
                if art_match:
                    markdown_lines.append(f"\n### {text}\n")
                else:
                    formatted = []
                    for run in paragraph.runs:
                        run_text = run.text
                        if run.bold:
                            run_text = f"**{run_text}**"
                        if run.italic:
                            run_text = f"*{run_text}*"
                        formatted.append(run_text)
                    markdown_lines.append("".join(formatted))

        for table in doc.tables:
            markdown_lines.append("\n")
            for i, row in enumerate(table.rows):
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                markdown_lines.append("| " + " | ".join(row_cells) + " |")
                if i == 0:
                    markdown_lines.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
            markdown_lines.append("\n")

        content = "\n".join(markdown_lines)
        content = re.sub(r'\n\s*\n', '\n\n', content)
        return content, cls._extract_sections(content)

    @classmethod
    def _parse_pdf_file(cls, file_path: str) -> tuple[str, list[dict]]:
        reader = pypdf.PdfReader(file_path)
        markdown_lines = []
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                markdown_lines.append(f"\n<!-- Page {page_num + 1} -->\n")
                markdown_lines.append(text)

        content = "\n".join(markdown_lines)
        content = re.sub(r'(\w)-\n(\w)', r'\1\2', content)
        return content, cls._extract_sections(content)

    @classmethod
    def _extract_sections(cls, markdown_content: str) -> list[dict]:
        """
        Parses Markdown titles and structures content into hierarchy sections (Chapters, Articles, Preambles).
        """
        sections = []
        lines = markdown_content.splitlines()
        current_section = {
            "title": "Преамбула",
            "article_num": None,
            "level": 1,
            "content_lines": []
        }

        header_regex = re.compile(r'^(#{1,6})\s+(.*)$')
        article_num_regex = re.compile(r'статья\s+(\d+[\-\d]*)', re.IGNORECASE)

        for line in lines:
            match = header_regex.match(line)
            if match:
                if current_section["content_lines"] or current_section["title"] != "Преамбула":
                    current_section["content"] = "\n".join(current_section["content_lines"]).strip()
                    del current_section["content_lines"]
                    sections.append(current_section)

                level = len(match.group(1))
                title = match.group(2).strip()
                num_match = article_num_regex.search(title)
                art_num = num_match.group(1) if num_match else None

                current_section = {
                    "title": title,
                    "article_num": art_num,
                    "level": level,
                    "content_lines": []
                }
            else:
                current_section["content_lines"].append(line)

        # Append last section
        if current_section["content_lines"] or current_section["title"] != "Преамбула":
            current_section["content"] = "\n".join(current_section["content_lines"]).strip()
            del current_section["content_lines"]
            sections.append(current_section)

        return sections
