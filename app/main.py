import os
import io
import re
import logging
import datetime
import warnings
warnings.filterwarnings("ignore", message=".*ARC4 has been moved to cryptography.*")
from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File, Form, BackgroundTasks, status
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from sqlalchemy import text
from pydantic import BaseModel, Field, ConfigDict
import docx
from google import genai

from app.database import get_db, init_db, SessionLocal
from app.models import Source, SourceVersion, InternalDoc, Task, Prompt, PromptHistory, Memory, User, Notification
from app.auth import get_current_user, create_access_token, verify_password
from fastapi.security import OAuth2PasswordRequestForm
from app.agents.monitoring import MonitoringAgent
from app.agents.legal import LegalAgent
from app.agents.document import DocumentAgent
from app.agents.knowledge import KnowledgeAgent
from app.agents.parsing import ParsingAgent
from app.agents.notifications import NotificationService
from app.scheduler import scheduler

from contextlib import asynccontextmanager

logger = logging.getLogger(__name__)


def _run_change_pipeline(task_id: int):
    """
    Background worker: runs legal analysis + draft/review generation for a
    freshly detected law change.
    """
    db = SessionLocal()
    try:
        task = db.query(Task).filter(Task.id == task_id).first()
        if not task:
            return
        old_v = task.old_version.content_markdown if task.old_version else ""
        new_v = task.new_version.content_markdown if task.new_version else ""

        analysis = LegalAgent.analyze_changes(old_v, new_v, db)
        task.analysis_json = analysis
        db.commit()

        doc_res = DocumentAgent.generate_impact_and_drafts(task_id, db)
        
        # Check review quality gate and notify if attention is needed
        score = doc_res.get("score", 0)
        source_title = task.source.title if task.source else "НПА"
        if score < 80:
            NotificationService.create_notification(
                title=f"⚠️ Требуется доработка проектов ({score}%)",
                message=f"Для задачи #{task.id} ({source_title}) оценка соответствия ниже порога (80%). Рекомендуется ручная проверка методолога.",
                type="quality_warning",
                severity="high",
                link_url="#tasks",
                db=db
            )
        else:
            NotificationService.create_notification(
                title=f"✅ Проекты подготовлены ({score}%)",
                message=f"Пакет изменений по задаче #{task.id} ({source_title}) успешно сформирован и верифицирован комплаенс-агентом.",
                type="change_detected",
                severity="normal",
                link_url="#tasks",
                db=db
            )
    except Exception as e:
        logger.error(f"_run_change_pipeline failed for task {task_id}: {e}")
        try:
            if task and task.source:
                NotificationService.create_notification(
                    title=f"❌ Ошибка генерации проектов ({task.source.title})",
                    message=f"Произошел сбой при автоматической подготовке документов: {str(e)}",
                    type="error",
                    severity="high",
                    link_url="#tasks",
                    db=db
                )
        except Exception:
            pass
    finally:
        db.close()


def _run_audit_pipeline(task_id: int):
    """Background worker: general compliance audit for a source with no detected text changes."""
    db = SessionLocal()
    try:
        task = db.query(Task).filter(Task.id == task_id).first()
        if not task:
            return
        current_text = task.new_version.content_markdown if task.new_version else ""

        analysis = LegalAgent.audit_compliance(current_text, db)
        task.analysis_json = analysis
        db.commit()

        doc_res = DocumentAgent.generate_impact_and_drafts(task_id, db)
        score = doc_res.get("score", 0)
        source_title = task.source.title if task.source else "НПА"
        
        NotificationService.create_notification(
            title=f"📋 Плановая проверка завершена ({score}%)",
            message=f"Плановый аудит соответствия действующей редакции {source_title} завершен (Задача #{task.id}).",
            type="audit_completed",
            severity="normal",
            link_url="#tasks",
            db=db
        )
    except Exception as e:
        logger.error(f"_run_audit_pipeline failed for task {task_id}: {e}")
    finally:
        db.close()


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    # Start automated background monitoring scheduler
    scheduler.start(change_pipeline_fn=_run_change_pipeline, audit_pipeline_fn=_run_audit_pipeline)
    yield
    # Stop scheduler cleanly on shutdown
    await scheduler.stop()


app = FastAPI(title="Compliance Centras API", version="1.1.0", lifespan=lifespan)

raw_origins = os.getenv("ALLOWED_ORIGINS", "*")
allowed_origins = [o.strip() for o in raw_origins.split(",") if o.strip()]
if "*" in allowed_origins or not allowed_origins:
    allowed_origins = ["*"]

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True if allowed_origins != ["*"] else False,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.middleware("http")
async def add_security_headers(request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-XSS-Protection"] = "1; mode=block"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    # Ensure browsers don't cache stale JS/CSS during continuous deployments
    if request.url.path.endswith((".js", ".css", ".html")) or request.url.path == "/":
        response.headers["Cache-Control"] = "no-cache, must-revalidate"
    return response

# --- System Healthcheck (Production Endpoint) ---
@app.get("/healthz")
@app.get("/api/health")
def health_check(db: Session = Depends(get_db)):
    db_status = "healthy"
    try:
        db.execute(text("SELECT 1")).scalar()
    except Exception as e:
        db_status = f"unhealthy: {str(e)}"
        logger.error(f"Healthcheck database error: {e}")

    scheduler_status = scheduler.get_status()
    is_healthy = db_status == "healthy"

    return JSONResponse(
        status_code=200 if is_healthy else 503,
        content={
            "status": "healthy" if is_healthy else "degraded",
            "timestamp": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "version": "1.1.0",
            "database": db_status,
            "scheduler": scheduler_status,
            "llm_configured": bool(os.getenv("GEMINI_API_KEY"))
        }
    )

# Pydantic Schemas for validation
class SourceCreate(BaseModel):
    title: str = Field(..., max_length=255)
    url: str
    check_interval_hours: float = 24.0

class SourceUpdate(BaseModel):
    title: str
    url: str
    check_interval_hours: float
    is_active: bool

class InternalDocCreate(BaseModel):
    title: str = Field(..., max_length=255)
    category: str = Field(..., max_length=100)
    content_markdown: str

class InternalDocOut(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    title: str
    category: str
    content_markdown: str
    created_at: datetime.datetime

class TaskStatusUpdate(BaseModel):
    status: str  # Approved, Rejected
    reject_reason: str = None

class PromptUpdate(BaseModel):
    content: str

class ChatMessage(BaseModel):
    message: str

class NotificationOut(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    title: str
    message: str
    type: str
    severity: str
    link_url: str | None = None
    is_read: bool
    created_at: datetime.datetime

# ----------------- FastAPI Endpoints -----------------

@app.post("/api/login")
def login(form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == form_data.username).first()
    if not user or not verify_password(form_data.password, user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect email or password",
            headers={"WWW-Authenticate": "Bearer"}
        )
    access_token = create_access_token(data={"sub": user.email})
    return {"access_token": access_token, "token_type": "bearer"}


@app.get("/api/dashboard")
def get_dashboard(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    sources_count = db.query(Source).count()
    active_sources = db.query(Source).filter(Source.is_active == True).count()
    pending_tasks = db.query(Task).filter(Task.status == "Pending").count()
    completed_tasks = db.query(Task).filter(Task.status == "Approved").count()
    total_docs = db.query(InternalDoc).count()
    unread_notifications = NotificationService.get_unread_count(db)
    
    # Recent tasks
    recent_tasks = db.query(Task).order_by(Task.created_at.desc()).limit(5).all()
    recent_tasks_data = []
    for t in recent_tasks:
        recent_tasks_data.append({
            "id": t.id,
            "source_title": t.source.title if t.source else "НПА",
            "status": t.status,
            "created_at": t.created_at.strftime("%Y-%m-%d %H:%M:%S") if t.created_at else "",
            "score": (t.analysis_json or {}).get("review_report", {}).get("score", 0) if t.analysis_json else 0
        })
        
    return {
        "sources_count": sources_count,
        "active_sources": active_sources,
        "pending_tasks": pending_tasks,
        "completed_tasks": completed_tasks,
        "total_docs": total_docs,
        "unread_notifications": unread_notifications,
        "recent_tasks": recent_tasks_data
    }


# --- Notifications ---
@app.get("/api/notifications")
def list_notifications(
    unread_only: bool = False,
    limit: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    notifications = NotificationService.get_notifications(db, limit=limit, unread_only=unread_only)
    unread_count = NotificationService.get_unread_count(db)
    notifs_data = [
        {
            "id": n.id,
            "title": n.title,
            "message": n.message,
            "type": n.type,
            "severity": n.severity,
            "link_url": n.link_url,
            "is_read": n.is_read,
            "created_at": n.created_at.strftime("%Y-%m-%d %H:%M:%S") if n.created_at else ""
        }
        for n in notifications
    ]
    return {
        "unread_count": unread_count,
        "notifications": notifs_data,
        "items": notifs_data
    }

@app.put("/api/notifications/{id}/read")
@app.post("/api/notifications/{id}/read")
def mark_notification_read(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    success = NotificationService.mark_as_read(id, db)
    if not success:
        raise HTTPException(status_code=404, detail="Notification not found")
    return {"message": "Notification marked as read"}

@app.post("/api/notifications/read-all")
def mark_all_notifications_read(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    count = NotificationService.mark_all_as_read(db)
    return {"message": f"{count} notifications marked as read"}

@app.post("/api/notifications/test")
def trigger_test_notification(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    n = NotificationService.create_notification(
        title="🔔 Тестовое уведомление",
        message="Система мониторинга и автоматического оповещения Compliance Centras активна.",
        type="info",
        severity="normal",
        link_url="#monitoring",
        db=db
    )
    return {"status": "success", "id": n.id}

@app.post("/api/notifications/telegram-test")
def trigger_telegram_test(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    chat_id = os.getenv("TELEGRAM_CHAT_ID")
    if not token or not chat_id:
        return {
            "status": "warning",
            "message": "TELEGRAM_BOT_TOKEN или TELEGRAM_CHAT_ID не заданы в .env. Для отправки в Telegram укажите их в конфигурации."
        }
    n = NotificationService.create_notification(
        title="🚨 ТЕСТОВЫЙ КОМПЛАЕНС АЛЕРТ",
        message="Проверка интеграции Telegram Bot для оповещения руководства и методологов Centras Insurance.",
        type="change_detected",
        severity="critical",
        link_url="http://localhost:8000/#tasks",
        db=db
    )
    return {"status": "success", "message": "Тестовое уведомление успешно отправлено в Telegram!", "id": n.id}


# --- Scheduler ---
@app.get("/api/scheduler/status")
def get_scheduler_status(current_user: User = Depends(get_current_user)):
    return scheduler.get_status()

@app.post("/api/scheduler/run-now")
def run_scheduler_now(current_user: User = Depends(get_current_user)):
    summary = scheduler.check_overdue_sources()
    return {"status": "success", "summary": summary}

# --- Sources ---
@app.get("/api/sources")
def list_sources(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Source).all()

@app.post("/api/sources")
def create_source(source_data: SourceCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    source = Source(
        title=source_data.title,
        url=source_data.url,
        check_interval_hours=source_data.check_interval_hours,
        is_active=True
    )
    db.add(source)
    db.commit()
    db.refresh(source)

    # Establish initial baseline Version 1 immediately upon source creation
    try:
        MonitoringAgent.run_check(source.id, db)
    except Exception as e:
        logger.warning(f"Initial baseline check for source {source.id} deferred: {e}")

    return source

@app.get("/api/sources/{id}/export")
def export_source_docx(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    source = db.query(Source).filter(Source.id == id).first()
    if not source:
        raise HTTPException(status_code=404, detail="Source not found")

    latest_version = db.query(SourceVersion).filter(
        SourceVersion.source_id == id
    ).order_by(SourceVersion.version_num.desc()).first()

    doc = docx.Document()
    doc.add_heading(source.title, level=0)

    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Первоисточник: {source.url}\n").italic = True
    if latest_version:
        meta_p.add_run(f"Редакция: Версия #{latest_version.version_num}\n")
        meta_p.add_run(f"Контрольная сумма (SHA-256): {latest_version.sha256}\n")
        if latest_version.created_at:
            meta_p.add_run(f"Дата фиксации: {latest_version.created_at.strftime('%Y-%m-%d %H:%M:%S')}\n")

        doc.add_heading("Текст нормативного акта", level=1)
        for line in (latest_version.content_markdown or "").splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("- "):
                doc.add_paragraph(line_str[2:], style='List Bullet')
            else:
                doc.add_paragraph(line_str)
    else:
        doc.add_paragraph("Текст редакции еще не сохранен.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=law_source_{id}.docx"}
    )

@app.put("/api/sources/{id}")
def update_source(id: int, source_data: SourceUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    source = db.query(Source).filter(Source.id == id).first()
    if not source:
        raise HTTPException(status_code=404, detail="Source not found")
    source.title = source_data.title
    source.url = source_data.url
    source.check_interval_hours = source_data.check_interval_hours
    source.is_active = source_data.is_active
    db.commit()
    db.refresh(source)
    return source

@app.delete("/api/sources/{id}")
def delete_source(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    source = db.query(Source).filter(Source.id == id).first()
    if not source:
        raise HTTPException(status_code=404, detail="Source not found")
    db.delete(source)
    db.commit()
    return {"message": "Source deleted successfully"}

@app.post("/api/sources/{id}/find_url")
def find_source_url(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    source = db.query(Source).filter(Source.id == id).first()
    if not source:
        raise HTTPException(status_code=404, detail="Source not found")
    return LegalAgent.find_source_url(source.title)

@app.post("/api/sources/{id}/check")
def check_source(id: int, background_tasks: BackgroundTasks, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    res = MonitoringAgent.run_check(id, db)

    if res.get("status") == "changed" and res.get("task_id"):
        background_tasks.add_task(_run_change_pipeline, res["task_id"])
        
        NotificationService.create_notification(
            title=f"🚨 Изменения в {res.get('source')}",
            message=f"Обнаружена новая редакция (версия #{res.get('version')}). Запущен правовой анализ.",
            type="change_detected",
            severity="high",
            link_url="#tasks",
            db=db
        )

    elif res.get("status") == "baseline_established":
        NotificationService.create_notification(
            title=f"📌 Зафиксирован эталон: {res.get('source')}",
            message=f"Первоначальная версия (#{res.get('version', 1)}) успешно зафиксирована в базе мониторинга (SHA-256).",
            type="info",
            severity="normal",
            link_url="#monitoring",
            db=db
        )

    elif res.get("status") in ("unchanged", "cosmetic_only"):
        existing_pending = db.query(Task).filter(
            Task.source_id == id,
            Task.status == "Pending"
        ).first()

        if not existing_pending:
            latest_version = db.query(SourceVersion).filter(
                SourceVersion.source_id == id
            ).order_by(SourceVersion.version_num.desc()).first()

            if latest_version:
                audit_task = Task(
                    source_id=id,
                    old_version_id=latest_version.id,
                    new_version_id=latest_version.id,
                    task_type="audit",
                    diff_html=(
                        "<div style='padding:20px;text-align:center;'>"
                        "Текст закона не изменился. Ниже — результаты плановой проверки "
                        "соответствия текущей редакции.</div>"
                    ),
                    status="Pending"
                )
                db.add(audit_task)
                db.commit()
                db.refresh(audit_task)

                background_tasks.add_task(_run_audit_pipeline, audit_task.id)
                res["audit_task_id"] = audit_task.id


    return res

@app.post("/api/sources/test-simulation")
def run_live_test_simulation(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """
    Creates or modifies a live test law file (data/laws/test_live_law.txt) with a new incremented clause,
    registers it as a Source with 1-minute interval, and prepares it for automated live detection.
    """
    laws_dir = os.path.join(os.path.dirname(__file__), "..", "data", "laws")
    os.makedirs(laws_dir, exist_ok=True)
    test_file = os.path.join(laws_dir, "test_live_law.txt")
    
    current_time_str = datetime.datetime.now().strftime("%H:%M:%S")
    rand_amount = 1000000 + int(datetime.datetime.now().timestamp()) % 900000
    
    content = (
        f"# Закон РК 'О страховании и комплаенс-контроле' (Тестовый Live-источник)\n\n"
        f"Статья 5. Минимальный гарантийный фонд страховой организации.\n"
        f"1. Размер фонда устанавливается в размере {rand_amount:,} тенге.\n"
        f"2. Требование обновлено в реальном времени в {current_time_str}.\n"
        f"3. Страховые компании обязаны исполнить норматив незамедлительно.\n"
    )
    
    with open(test_file, "w", encoding="utf-8") as f:
        f.write(content)
        
    src = db.query(Source).filter(Source.title == "Тестовый закон (Live-симуляция)").first()
    if not src:
        src = Source(
            title="Тестовый закон (Live-симуляция)",
            url=test_file,
            check_interval_hours=0.0167,  # 1 minute
            is_active=True
        )
        db.add(src)
        db.commit()
        db.refresh(src)
    else:
        src.url = test_file
        src.check_interval_hours = 0.0167
        src.is_active = True
        db.commit()
        
    return {
        "status": "success",
        "message": f"Тестовый закон обновлен в {current_time_str} (Фонд: {rand_amount:,} ₸). Интервал: 1 мин.",
        "source_id": src.id
    }

# --- Internal Documents ---
@app.get("/api/documents", response_model=list[InternalDocOut])
def list_documents(search: str = Query(None), db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query = db.query(InternalDoc)
    if search:
        query = query.filter(
            (InternalDoc.title.like(f"%{search}%")) |
            (InternalDoc.content_markdown.like(f"%{search}%"))
        )
    return query.all()

@app.post("/api/documents", response_model=InternalDocOut)
def create_document(doc_data: InternalDocCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = KnowledgeAgent.register_document(
        title=doc_data.title,
        category=doc_data.category,
        content=doc_data.content_markdown,
        db=db
    )
    return doc

@app.post("/api/documents/upload", response_model=InternalDocOut)
async def upload_document(
    title: str = Form(...),
    category: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    content_bytes = await file.read()
    safe_name = os.path.basename(file.filename or "upload")
    try:
        markdown_content, _structure = ParsingAgent.parse_content(content_bytes, safe_name)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Не удалось обработать файл: {e}")

    if not markdown_content.strip():
        raise HTTPException(status_code=400, detail="Не удалось извлечь текст из файла.")

    doc = KnowledgeAgent.register_document(
        title=title,
        category=category,
        content=markdown_content,
        db=db
    )
    return doc

@app.get("/api/documents/{id}", response_model=InternalDocOut)
def get_document(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = db.query(InternalDoc).filter(InternalDoc.id == id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    return doc

@app.get("/api/documents/{id}/export")
def export_internal_doc_docx(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    internal_doc = db.query(InternalDoc).filter(InternalDoc.id == id).first()
    if not internal_doc:
        raise HTTPException(status_code=404, detail="Document not found")

    doc = docx.Document()
    doc.add_heading(internal_doc.title, level=0)

    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Категория: {internal_doc.category}\n").italic = True
    if internal_doc.created_at:
        meta_p.add_run(f"Дата регистрации: {internal_doc.created_at.strftime('%Y-%m-%d %H:%M:%S')}\n")

    doc.add_heading("Содержание регламента", level=1)
    for line in (internal_doc.content_markdown or "").splitlines():
        line_str = line.strip()
        if not line_str:
            continue
        if line_str.startswith("# "):
            doc.add_heading(line_str[2:], level=1)
        elif line_str.startswith("## "):
            doc.add_heading(line_str[3:], level=2)
        elif line_str.startswith("### "):
            doc.add_heading(line_str[4:], level=3)
        elif line_str.startswith("- "):
            doc.add_paragraph(line_str[2:], style='List Bullet')
        else:
            doc.add_paragraph(line_str)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=internal_doc_{id}.docx"}
    )

@app.post("/api/documents/{id}/suggest_sources")
def suggest_sources_for_document(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = db.query(InternalDoc).filter(InternalDoc.id == id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    suggestions, error = LegalAgent.suggest_monitoring_sources(doc.content_markdown, db)
    return {"suggestions": suggestions, "error": error}

@app.delete("/api/documents/{id}")
def delete_document(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = db.query(InternalDoc).filter(InternalDoc.id == id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    db.delete(doc)
    db.commit()
    return {"message": "Document deleted successfully"}

# --- Tasks ---
@app.get("/api/tasks")
def list_tasks(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    tasks = db.query(Task).order_by(Task.created_at.desc()).all()
    results = []
    for t in tasks:
        title = t.source.title if t.source else "НПА"
        if t.task_type == "audit":
            title += " (плановая проверка соответствия)"
        results.append({
            "id": t.id,
            "source_title": title,
            "status": t.status,
            "created_at": t.created_at.strftime("%Y-%m-%d %H:%M:%S") if t.created_at else "",
            "has_drafts": t.draft_json is not None
        })
    return results

@app.get("/api/tasks/{id}")
def get_task(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")
        
    related_docs = []
    if t.analysis_json:
        query_text = f"{t.analysis_json.get('what_changed', '')} {t.analysis_json.get('requirements', '')}"
        similar = KnowledgeAgent.search_similar_docs(query_text, db, limit=4)
        for doc, score in similar:
            related_docs.append({
                "id": doc.id,
                "title": doc.title,
                "category": doc.category,
                "similarity": round(score, 2)
            })
            
    return {
        "id": t.id,
        "source_title": t.source.title if t.source else "НПА",
        "source_url": t.source.url if t.source else "",
        "old_version_num": t.old_version.version_num if t.old_version else None,
        "new_version_num": t.new_version.version_num if t.new_version else None,
        "status": t.status,
        "reject_reason": t.reject_reason,
        "diff_html": t.diff_html,
        "analysis": t.analysis_json,
        "drafts": t.draft_json,
        "related_docs": related_docs,
        "created_at": t.created_at.strftime("%Y-%m-%d %H:%M:%S") if t.created_at else ""
    }

@app.put("/api/tasks/{id}")
def update_task_status(id: int, data: TaskStatusUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")
        
    t.status = data.status
    source_title = t.source.title if t.source else "НПА"
    if data.status == "Rejected":
        t.reject_reason = data.reject_reason
        
        # Store failed decision reason in Corporate Memory
        memory_content = f"Решение методолога отклонено по задаче #{t.id} ({source_title}). Причина: {data.reject_reason}."
        mem = Memory(
            type="decision",
            content_text=memory_content,
            metadata_json={"task_id": t.id, "source": source_title, "reject_reason": data.reject_reason}
        )
        db.add(mem)
    else:
        t.reject_reason = None
        # Store accepted decision context in Memory
        memory_content = f"Успешное согласование изменений методологии по задаче #{t.id} ({source_title})."
        mem = Memory(
            type="decision",
            content_text=memory_content,
            metadata_json={"task_id": t.id, "source": source_title}
        )
        db.add(mem)
        
    db.commit()
    return {"message": f"Task status updated to {data.status}"}

@app.post("/api/tasks/{id}/regenerate")
def regenerate_task_drafts(id: int, background_tasks: BackgroundTasks, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")

    t.draft_json = None
    db.commit()

    if t.task_type == "audit":
        background_tasks.add_task(_run_audit_pipeline, t.id)
    else:
        background_tasks.add_task(_run_change_pipeline, t.id)

    return {"status": "queued"}

@app.post("/api/tasks/{id}/save_drafts")
def save_task_drafts(id: int, data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")
    t.draft_json = data
    db.commit()
    return {"message": "Drafts saved successfully"}

# --- Document Exporting ---
@app.get("/api/tasks/{id}/export/{doc_type}")
def export_docx(id: int, doc_type: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")
        
    drafts = t.draft_json or {}
    
    doc = docx.Document()
    doc.add_heading(f"Compliance Centras - Экспорт", 0)
    
    if doc_type == "memo":
        doc.add_heading("Служебная записка", 1)
        doc.add_paragraph(str(drafts.get("memo", "Текст отсутствует.")))
    elif doc_type == "rules":
        doc.add_heading("Проект изменений во внутренние правила", 1)
        doc.add_paragraph(str(drafts.get("rules_changes", "Текст отсутствует.")))
    elif doc_type == "board_letter":
        doc.add_heading("Сопроводительное письмо Правлению", 1)
        doc.add_paragraph(str(drafts.get("board_letter", "Текст отсутствует.")))
    elif doc_type == "board_directors_letter":
        doc.add_heading("Сопроводительное письмо Совету директоров", 1)
        doc.add_paragraph(str(drafts.get("board_directors_letter", "Текст отсутствует.")))
    elif doc_type == "plan":
        doc.add_heading("План внедрения изменений", 1)
        plan = drafts.get("implementation_plan", [])
        if isinstance(plan, list):
            for item in plan:
                if isinstance(item, dict):
                    doc.add_paragraph(f"- Этап: {item.get('step', '')}\n  Срок: {item.get('deadline', '')}\n  Ответственный: {item.get('owner', '')}")
                else:
                    doc.add_paragraph(f"- {str(item)}")
        else:
            doc.add_paragraph(str(plan))
    elif doc_type == "checklist":
        doc.add_heading("Комплаенс Чек-лист", 1)
        checklist = drafts.get("checklist", [])
        if isinstance(checklist, list):
            for item in checklist:
                if isinstance(item, dict):
                    doc.add_paragraph(f"[ ] {item.get('item', '')}")
                else:
                    doc.add_paragraph(f"[ ] {str(item)}")
        else:
            doc.add_paragraph(str(checklist))
    else:
        raise HTTPException(status_code=400, detail="Invalid document type")
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=compliance_{doc_type}_{t.id}.docx"}
    )

@app.get("/api/tasks/{id}/export_table")
def export_task_comparative_table_docx(id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    t = db.query(Task).filter(Task.id == id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Task not found")
    try:
        docx_bytes = MonitoringAgent.generate_comparative_table_docx(t.id, db)
        return StreamingResponse(
            io.BytesIO(docx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=comparative_table_task_{t.id}.docx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate comparative table: {e}")

# --- Relations Graph ---
@app.get("/api/relations")
def get_relations_graph(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    nodes = []
    links = []
    node_ids = set()

    def add_node(node_id: str, label: str, node_type: str, details: dict = None):
        if node_id not in node_ids:
            node_ids.add(node_id)
            nodes.append({"id": node_id, "label": label, "type": node_type, "details": details or {}})

    for source in db.query(Source).all():
        law_id = f"law-{source.id}"
        add_node(law_id, source.title, "law", {"url": source.url, "interval": source.check_interval_hours})

        tasks = db.query(Task).filter(
            Task.source_id == source.id, Task.analysis_json.isnot(None)
        ).order_by(Task.created_at.desc()).all()

        for task in tasks:
            task_id = f"task-{task.id}"
            add_node(task_id, f"Задача #{task.id} ({task.status})", "task", {"status": task.status, "task_id": task.id})
            links.append({"source": law_id, "target": task_id})

            analysis = task.analysis_json or {}
            query_text = f"{analysis.get('what_changed', '')} {analysis.get('requirements', '')}".strip()
            if query_text:
                for doc, score in KnowledgeAgent.search_similar_docs(query_text, db, limit=2):
                    doc_id = f"doc-{doc.id}"
                    add_node(doc_id, doc.title, "rule", {"category": doc.category, "doc_id": doc.id, "similarity": round(score, 2)})
                    links.append({"source": task_id, "target": doc_id})

                    # If drafts contain specific rules changes, create a rule clause node
                    drafts = task.draft_json or {}
                    if drafts.get("rules_changes"):
                        clause_id = f"clause-{task.id}-{doc.id}"
                        add_node(clause_id, f"Правки: {doc.title[:22]}...", "clause", {
                            "summary": str(drafts.get("rules_changes", ""))[:140]
                        })
                        links.append({"source": doc_id, "target": clause_id})

            for dept in analysis.get("affected_departments", []):
                dept_name = ""
                if isinstance(dept, dict):
                    dept_name = (dept.get("department") or "").strip()
                elif isinstance(dept, str):
                    dept_name = dept.strip()
                if dept_name:
                    dept_id = f"dept-{dept_name}"
                    add_node(dept_id, dept_name, "dept", {"department": dept_name})
                    links.append({"source": task_id, "target": dept_id})

    # Include all internal documents in the catalog
    for internal_doc in db.query(InternalDoc).all():
        doc_id = f"doc-{internal_doc.id}"
        add_node(doc_id, internal_doc.title, "rule", {"category": internal_doc.category, "doc_id": internal_doc.id})

    return {"nodes": nodes, "links": links}

# --- Prompt Zone ---
@app.get("/api/prompts")
def get_prompts(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Prompt).all()

@app.put("/api/prompts/{key}")
def update_prompt(key: str, data: PromptUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    prompt = db.query(Prompt).filter(Prompt.key == key).first()
    if not prompt:
        raise HTTPException(status_code=404, detail="Prompt key not found")
        
    new_version = prompt.version + 1
    
    history = PromptHistory(
        prompt_id=prompt.id,
        key=prompt.key,
        content=prompt.content,
        version=prompt.version
    )
    db.add(history)
    
    prompt.content = data.content
    prompt.version = new_version
    db.commit()
    db.refresh(prompt)
    return prompt

# --- AI Chat ---
@app.post("/api/chat")
def handle_chat(chat_data: ChatMessage, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    similar = KnowledgeAgent.search_similar_docs(chat_data.message, db, limit=2)
    context_str = ""
    for doc, sim in similar:
        context_str += f"\nВнутренний документ: {doc.title}\n{doc.content_markdown[:800]}\n---\n"
        
    model_name = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return {"response": "AI Методолог готов к работе. Для генерации ответов на основе нейросети укажите GEMINI_API_KEY в настройках (.env)."}
    
    prompt = (
        "Вы — Старший методолог страховой компании Республики Казахстан. "
        "Ответьте на вопрос пользователя, используя предоставленный контекст внутренних документов и ваши знания законов РК.\n\n"
        f"КОНТЕКСТ:\n{context_str}\n\n"
        f"ВОПРОС: {chat_data.message}\n\n"
        "Ответьте развернуто, структурировано и на русском языке."
    )
    
    client = genai.Client(api_key=api_key)
    try:
        response = client.models.generate_content(
            model=model_name,
            contents=prompt
        )
        return {"response": response.text.strip()}
    except Exception as e:
        logger.error(f"Error in handle_chat Gemini call: {e}")
        return {"response": f"Произошла ошибка при обращении к ИИ: {str(e)}"}

class SandboxCompareRequest(BaseModel):
    old_text: str
    new_text: str


class SandboxUrlRequest(BaseModel):
    url: str

@app.post("/api/sandbox/parse-file")
async def sandbox_parse_file(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user)
):
    """Parses an uploaded PDF, Word DOCX, TXT or Markdown file into normalized text for Sandbox."""
    content_bytes = await file.read()
    safe_name = os.path.basename(file.filename or "uploaded_doc")
    try:
        markdown_text, _structure = ParsingAgent.parse_content(content_bytes, safe_name)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка парсинга файла: {e}")
        
    if not markdown_text.strip():
        raise HTTPException(status_code=400, detail="Не удалось извлечь текст из документа.")
        
    return {
        "text": markdown_text,
        "filename": safe_name,
        "chars": len(markdown_text),
        "words": len(markdown_text.split()),
        "lines": len(markdown_text.splitlines())
    }

@app.post("/api/sandbox/fetch-url")
def sandbox_fetch_url(
    req: SandboxUrlRequest,
    current_user: User = Depends(get_current_user)
):
    """Fetches text from a web link (Zan.kz, Adilet, etc.) or local file for Sandbox."""
    url = req.url.strip()
    if not url:
        raise HTTPException(status_code=400, detail="Укажите корректный URL или путь.")
    try:
        content_bytes = MonitoringAgent.fetch_source_content(url)
        clean_url = url.replace('\\', '/').split('?')[0]
        file_name = os.path.basename(clean_url) if clean_url else "doc.html"
        if '.' not in file_name:
            file_name += ".html"
        markdown_text, _ = ParsingAgent.parse_content(content_bytes, file_name)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Не удалось загрузить документ по ссылке: {e}")

    if not markdown_text.strip():
        raise HTTPException(status_code=400, detail="Текст документа по указанной ссылке пуст.")

    return {
        "text": markdown_text,
        "url": url,
        "chars": len(markdown_text),
        "words": len(markdown_text.split()),
        "lines": len(markdown_text.splitlines())
    }

@app.post("/api/sandbox/compare")
def sandbox_compare(
    req: SandboxCompareRequest,
    current_user: User = Depends(get_current_user)
):
    """
    Test comparison laboratory: analyzes differences between two text versions,
    calculates hashes, word/character deltas, cosmetic filter status, and renders side-by-side diff.
    """
    sha_old = MonitoringAgent.calculate_sha256(req.old_text)
    sha_new = MonitoringAgent.calculate_sha256(req.new_text)
    has_real_changes, diff_html = MonitoringAgent.diff_documents(req.old_text, req.new_text)
    
    return {
        "is_identical": sha_old == sha_new,
        "has_real_changes": has_real_changes,
        "sha256_old": sha_old,
        "sha256_new": sha_new,
        "chars_old": len(req.old_text),
        "chars_new": len(req.new_text),
        "chars_delta": len(req.new_text) - len(req.old_text),
        "words_old": len(req.old_text.split()),
        "words_new": len(req.new_text.split()),
        "lines_old": len(req.old_text.splitlines()),
        "lines_new": len(req.new_text.splitlines()),
        "diff_html": diff_html
    }

@app.post("/api/sandbox/analyze")
def sandbox_analyze(
    req: SandboxCompareRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Test sandbox AI legal analysis: sends both versions to Gemini to generate structured compliance findings.
    """
    analysis = LegalAgent.analyze_changes(req.old_text, req.new_text, db)
    return {"analysis": analysis}





# Mount frontend static directory
static_dir = os.path.join(os.path.dirname(__file__), "..", "static")
if os.path.exists(static_dir):
    app.mount("/", StaticFiles(directory=static_dir, html=True), name="static")
